import pytest
from openpyxl import load_workbook

from src.schemas.vacancy import VacancyDTO
from src.utils.exports import export_vacancies_to_csv, export_vacancies_to_xlsx

MOCK_VACANCIES = [
    VacancyDTO(
        title="Python Developer",
        company_name="Cyber Core Tech",
        salary="от 150 000 до 250 000 ₽",
        link="https://habr.com",
    ),
    VacancyDTO(
        title="FastAPI Engineer",
        company_name="Async Team",
        salary="300 000 ₽",
        link="https://hh.ru",
    ),
]


@pytest.mark.asyncio
async def test_export_vacancies_to_csv_success(tmp_path):
    test_file = tmp_path / "test_report.csv"

    await export_vacancies_to_csv(MOCK_VACANCIES, file_path=str(test_file))

    assert test_file.exists()

    rows = test_file.read_text(encoding="utf-8").splitlines()

    assert rows[0] == "Title,Company_name,Salary,Link"
    assert (
        rows[1]
        == "Python Developer,Cyber Core Tech,от 150 000 до 250 000 ₽,https://habr.com"
    )


@pytest.mark.asyncio
async def test_export_vacancies_to_csv_empty(tmp_path):
    vacancies = []

    test_file = tmp_path / "test_report.csv"

    result = await export_vacancies_to_csv(vacancies, file_path=str(test_file))

    assert not test_file.exists()
    assert result is None


@pytest.mark.asyncio
async def test_export_vacancies_to_csv_os_error(tmp_path):
    test_file = tmp_path / "invalid_dir\0_123/inside/report.csv"

    result = await export_vacancies_to_csv(MOCK_VACANCIES, file_path=str(test_file))

    assert not test_file.exists()
    assert result is None


@pytest.mark.asyncio
async def test_export_vacancies_to_xlsx_success(tmp_path):
    test_file = tmp_path / "test_report.xlsx"

    await export_vacancies_to_xlsx(MOCK_VACANCIES, file_path=str(test_file))

    assert test_file.exists()

    wb = load_workbook(str(test_file))
    ws = wb.active

    assert ws.title == "Vacancies"

    assert ws.cell(row=1, column=1).value == "Title"
    assert ws.cell(row=1, column=2).value == "Company_name"
    assert ws.cell(row=1, column=3).value == "Salary"
    assert ws.cell(row=1, column=4).value == "Link"

    assert ws.cell(row=2, column=1).value == "Python Developer"
    assert ws.cell(row=2, column=2).value == "Cyber Core Tech"
    assert ws.cell(row=2, column=3).value == "от 150 000 до 250 000 ₽"
    assert ws.cell(row=2, column=4).value == "https://habr.com"


@pytest.mark.asyncio
async def test_export_vacancies_to_xlsx_empty(tmp_path):
    vacancies = []

    test_file = tmp_path / "test_report.xlsx"

    result = await export_vacancies_to_xlsx(vacancies, file_path=str(test_file))

    assert not test_file.exists()
    assert result is None


@pytest.mark.asyncio
async def test_export_vacancies_to_xlsx_os_error(tmp_path):
    test_file = tmp_path / "invalid_dir\0_123/inside/report.xlsx"

    result = await export_vacancies_to_xlsx(MOCK_VACANCIES, file_path=str(test_file))

    assert not test_file.exists()
    assert result is None
